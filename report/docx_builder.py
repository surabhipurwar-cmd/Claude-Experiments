"""
DOCX report builder — generates a formatted Word document from enriched deal
data and AI trend analysis. Designed to be easily editable after generation.
"""

import os
from datetime import date
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _heading(doc: Document, text: str, level: int) -> None:
    doc.add_heading(text, level=level)


def _bullet(doc: Document, text: str, level: int = 0) -> None:
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.25 * (level + 1))
    p.add_run(text)


def _labelled_para(doc: Document, label: str, value: str) -> None:
    p = doc.add_paragraph()
    run_label = p.add_run(f"{label}: ")
    run_label.bold = True
    p.add_run(value or "—")


def _team_section(doc: Document, team_name: str, deals: list[dict], highlight: str) -> None:
    _heading(doc, f"{team_name} Deals", level=2)
    if highlight:
        p = doc.add_paragraph()
        p.add_run("AI Highlights: ").bold = True
        p.add_run(highlight)
        doc.add_paragraph()

    if not deals:
        doc.add_paragraph("No deals assigned to this team today.")
        return

    for deal in deals:
        merchant = (
            deal.get("Merchant Name") or deal.get("Account Name")
            or deal.get("Merchant") or deal.get("Account") or "Unknown Merchant"
        )
        _heading(doc, merchant, level=3)
        _labelled_para(doc, "Summary", deal.get("summary", ""))

        if deal.get("rep_context"):
            _labelled_para(doc, "Rep Context", deal["rep_context"])

        if deal.get("risk_flags"):
            p = doc.add_paragraph()
            p.add_run("Risk Flags:").bold = True
            for flag in deal["risk_flags"]:
                _bullet(doc, flag)

        if deal.get("opportunities"):
            p = doc.add_paragraph()
            p.add_run("Opportunities:").bold = True
            for opp in deal["opportunities"]:
                _bullet(doc, opp)

        if deal.get("chorus_calls"):
            p = doc.add_paragraph()
            p.add_run("Chorus Call Evidence:").bold = True
            for call in deal["chorus_calls"]:
                title = call.get("title") or call.get("call_id") or "Call"
                call_date = call.get("date", "")
                summary = call.get("summary") or "(no summary available)"
                _bullet(doc, f"{title} ({call_date}): {summary}")

        doc.add_paragraph()  # spacing between deals


# ---------------------------------------------------------------------------
# Main builder
# ---------------------------------------------------------------------------

def build_report(
    enriched_deals: list[dict],
    trends: dict,
    report_date: date,
    output_path: str,
) -> str:
    """
    Build a .docx report and save to output_path.
    Returns the final file path.
    """
    doc = Document()

    # --- Cover / title ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(f"Deal Desk Daily Report")
    run.bold = True
    run.font.size = Pt(24)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.add_run(report_date.strftime("%B %d, %Y")).font.size = Pt(14)

    stats = doc.add_paragraph()
    stats.alignment = WD_ALIGN_PARAGRAPH.CENTER
    total = len(enriched_deals)
    oam_n = sum(1 for d in enriched_deals if "OAM" in str(d.get("assigned_to", "")).upper())
    iam_n = sum(1 for d in enriched_deals if "IAM" in str(d.get("assigned_to", "")).upper())
    bd_n  = sum(1 for d in enriched_deals if "BD"  in str(d.get("assigned_to", "")).upper())
    stats.add_run(f"Total Deals: {total}  |  OAM: {oam_n}  |  IAM: {iam_n}  |  BD: {bd_n}")

    doc.add_page_break()

    # --- Executive Summary ---
    _heading(doc, "Executive Summary", level=1)
    doc.add_paragraph(trends.get("executive_summary", "No summary generated."))
    doc.add_paragraph()

    # --- Top Trends ---
    _heading(doc, "Top Trends", level=1)
    top_trends = trends.get("top_trends", [])
    if top_trends:
        for trend in top_trends:
            p = doc.add_paragraph()
            p.add_run(trend.get("title", "")).bold = True
            doc.add_paragraph(trend.get("detail", ""))
    else:
        doc.add_paragraph("No trends identified.")
    doc.add_paragraph()

    # --- Red Flags ---
    _heading(doc, "Red Flags", level=1)
    red_flags = trends.get("red_flags", [])
    if red_flags:
        for flag in red_flags:
            _bullet(doc, flag)
    else:
        doc.add_paragraph("None today.")
    doc.add_paragraph()

    # --- Wins ---
    _heading(doc, "Wins & Bright Spots", level=1)
    wins = trends.get("wins", [])
    if wins:
        for win in wins:
            _bullet(doc, win)
    else:
        doc.add_paragraph("None today.")
    doc.add_paragraph()

    # --- Recommended Actions ---
    _heading(doc, "Recommended Actions for Leadership", level=1)
    actions = trends.get("recommended_actions", [])
    if actions:
        for action in actions:
            _bullet(doc, action)
    doc.add_paragraph()

    doc.add_page_break()

    # --- Deal Detail by Team ---
    _heading(doc, "Deal Detail by Team", level=1)

    oam_deals = [d for d in enriched_deals if "OAM" in str(d.get("assigned_to", "")).upper()]
    iam_deals = [d for d in enriched_deals if "IAM" in str(d.get("assigned_to", "")).upper()]
    bd_deals  = [d for d in enriched_deals if "BD"  in str(d.get("assigned_to", "")).upper()]
    other     = [d for d in enriched_deals if d not in oam_deals + iam_deals + bd_deals]

    _team_section(doc, "OAM (Outside Account Manager)", oam_deals, trends.get("oam_highlights", ""))
    doc.add_page_break()
    _team_section(doc, "IAM (Inside Account Manager)", iam_deals, trends.get("iam_highlights", ""))
    doc.add_page_break()
    _team_section(doc, "BD (Renegotiation Specialists)", bd_deals, trends.get("bd_highlights", ""))

    if other:
        doc.add_page_break()
        _team_section(doc, "Unassigned / Other", other, "")

    # --- Save ---
    os.makedirs(os.path.dirname(output_path) if os.path.dirname(output_path) else ".", exist_ok=True)
    doc.save(output_path)
    print(f"[docx] Report saved to {output_path}")
    return output_path
